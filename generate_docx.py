import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os

def create_document():
    doc = Document()

    # Configure Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Fonts
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Helper for adding formatted code block
    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)

        # Set cell background color to light gray
        shading = parse_xml(r'<w:shd {} w:fill="F5F5F5"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)

        # Borders
        borders = parse_xml(
            r'<w:tcBorders {}><w:top w:val="single" w:sz="4" w:space="0" w:color="DDDDDD"/>'
            r'<w:left w:val="single" w:sz="24" w:space="0" w:color="0066CC"/>'
            r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="DDDDDD"/>'
            r'<w:right w:val="single" w:sz="4" w:space="0" w:color="DDDDDD"/></w:tcBorders>'.format(nsdecls('w'))
        )
        cell._tc.get_or_add_tcPr().append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        run = p.add_run(code_text.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

        # Add spacing after table
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("LAPORAN PENJELASAN KODE & ARSITEKTUR\nAPLIKASI KALKULATOR ANDROID")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)
    p_title.paragraph_format.space_after = Pt(4)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Berbasis Kotlin, Jetpack Compose, & Material Design 3\nMahasiswa Teknik Informatika Semester 5")
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p_sub.paragraph_format.space_after = Pt(24)

    # Section 1: Pendahuluan
    h1 = doc.add_heading("1. Pendahuluan & Arsitektur Aplikasi", level=1)
    h1.style.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)

    doc.add_paragraph(
        "Aplikasi Kalkulator ini dikembangkan menggunakan bahasa pemrograman Kotlin modern dan framework UI deklaratif "
        "Jetpack Compose dengan skema desain Material Design 3. Aplikasi menerapkan arsitektur MVVM (Model-View-ViewModel) "
        "untuk memisahkan logika perhitungan (Business Logic) dari antarmuka pengguna (UI Layer)."
    )

    doc.add_paragraph("Struktur direktori proyek disusun sebagai berikut:")

    add_code_block(
"""com.example.kalkulator
├── MainActivity.kt                          <- Entry Point Utama Aplikasi
├── logic/
│   └── CalculatorViewModel.kt               <- Pengelolaan State & Logika Matematika
└── ui/
    ├── components/
    │   └── CalculatorButton.kt              <- Komponen Reusable Tombol UI
    ├── screens/
    │   └── CalculatorScreen.kt              <- Layout Utama Layar Kalkulator & Dialog Riwayat
    └── theme/
        ├── Color.kt                         <- Skema Warna Tema
        ├── Theme.kt                         <- Konfigurasi Material 3 Theme
        └── Type.kt                          <- Tipografi Font"""
    )

    # Section 2: Kode Lengkap Berkomentar
    h2 = doc.add_heading("2. Kode Lengkap & Komentar Penjelasan", level=1)
    h2.style.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)

    # 2.1 CalculatorViewModel.kt
    doc.add_heading("2.1 File Logika & State Management: CalculatorViewModel.kt", level=2)
    doc.add_paragraph(
        "File ini bertindak sebagai ViewModel yang menyimpan seluruh status (state) kalkulator "
        "dan mengeksekusi operasi matematika dasar maupun berantai sesuai aturan Kabataku."
    )

    vm_code = open("app/src/main/java/com/example/kalkulator/logic/CalculatorViewModel.kt", "r", encoding="utf-8").read()
    add_code_block(vm_code)

    # 2.2 CalculatorButton.kt
    doc.add_heading("2.2 File Reusable Component: CalculatorButton.kt", level=2)
    doc.add_paragraph(
        "File ini berisi komponen tombol kustom (Reusable Composable) yang digunakan kembali untuk seluruh tombol kalkulator."
    )
    btn_code = open("app/src/main/java/com/example/kalkulator/ui/components/CalculatorButton.kt", "r", encoding="utf-8").read()
    add_code_block(btn_code)

    # 2.3 CalculatorScreen.kt
    doc.add_heading("2.3 File Layout Utama: CalculatorScreen.kt", level=2)
    doc.add_paragraph(
        "File ini mengatur tata letak antarmuka (UI) layar kalkulator, tampilan ekspresi, dan dialog riwayat perhitungan."
    )
    screen_code = open("app/src/main/java/com/example/kalkulator/ui/screens/CalculatorScreen.kt", "r", encoding="utf-8").read()
    add_code_block(screen_code)

    # 2.4 MainActivity.kt
    doc.add_heading("2.4 File Entry Point: MainActivity.kt", level=2)
    doc.add_paragraph(
        "File Activity utama yang memuat konteks aplikasi dan menampilkan Composable CalculatorScreen."
    )
    main_code = open("app/src/main/java/com/example/kalkulator/MainActivity.kt", "r", encoding="utf-8").read()
    add_code_block(main_code)

    # Section 3: Alur Kerja Aplikasi
    h3 = doc.add_heading("3. Alur Kerja Aplikasi (Data & UI Flow)", level=1)
    h3.style.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)

    flow_p = doc.add_paragraph()
    flow_p.add_run("1. Sentuhan Layar / Event Input Pengguna (UI Layer):\n").bold = True
    flow_p.add_run("   Pengguna menekan tombol di layar (misalnya angka '5' atau operator '+'). Komponen CalculatorButton merespons event sentuhan dan memicu lambda callback onClick.\n\n")
    flow_p.add_run("2. Pengiriman Event ke ViewModel (Listener / Action):\n").bold = True
    flow_p.add_run("   CalculatorScreen memanggil fungsi listener pada CalculatorViewModel (seperti onNumberClick('5'), onOperatorClick('+'), atau onEqualsClick()).\n\n")
    flow_p.add_run("3. Pemrosesan Logika Bisnis & Pendeteksian Edge Case (ViewModel Layer):\n").bold = True
    flow_p.add_run("   CalculatorViewModel mengolah token angka/operator. Saat tombol '=' ditekan, fungsi evaluateTokens() mengeksekusi kalkulasi berdasarkan hierarki Kabataku (Perkalian/Pembagian dahulu, lalu Penjumlahan/Pengurangan) serta memeriksa pencegahan error pembagian dengan 0.\n\n")
    flow_p.add_run("4. Pembaharuan State (State Management):\n").bold = True
    flow_p.add_run("   ViewModel memperbarui variabel mutableStateOf (inputDisplay dan expressionDisplay) serta menambahkan item baru ke historyList.\n\n")
    flow_p.add_run("5. Penggambaran Ulang Layar Otomatis (Recomposition):\n").bold = True
    flow_p.add_run("   Jetpack Compose mendeteksi perubahan nilai state dan memicu Recomposition secara otomatis untuk menampilkan teks atau hasil kalkulasi terbaru ke layar pengguna.")

    # Section 4: Penanganan Edge Case
    h4 = doc.add_heading("4. Penanganan Edge Cases & Fitur Keamanan", level=1)
    h4.style.font.color.rgb = RGBColor(0x00, 0x44, 0xCC)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Kondisi / Edge Case"
    hdr_cells[1].text = "Penanganan di Kode"
    hdr_cells[2].text = "Alasan Teknis & Manfaat"

    for cell in hdr_cells:
        shading = parse_xml(r'<w:shd {} w:fill="0044CC"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data = [
        ("Pembagian dengan 0", "if (op == '÷' && b == 0.0) return Pair(null, 'Tidak dapat dibagi 0')", "Mencegah nilai Infinity/NaN atau ArithmeticException yang dapat menyebabkan crash pada aplikasi."),
        ("Titik Desimal Ganda", "if (!tokens.last().contains('.'))", "Memastikan sebuah segmen angka hanya memiliki maksimal satu titik desimal agar parsing Double aman."),
        ("Pencegahan Nol di Depan", "if (lastToken == '0') tokens[lastIndex] = number", "Menghindari bug format angka matematika seperti '05' atau '007'."),
        ("Meluapnya Teks Display", "fontSize = when { length > 18 -> 20.sp ... }", "Ukuran font mengecil secara dinamis agar angka panjang tidak overflow keluar layar.")
    ]

    for condition, code, reason in data:
        row_cells = table.add_row().cells
        row_cells[0].text = condition
        row_cells[1].text = code
        row_cells[2].text = reason

    # Save to Desktop and Project Folder
    desktop_path = os.path.expanduser("~/Desktop/Laporan_Kalkulator_Android.docx")
    project_path = "Laporan_Kalkulator_Android.docx"

    doc.save(project_path)
    try:
        doc.save(desktop_path)
        print(f"Berhasil disimpan di Desktop: {desktop_path}")
    except Exception as e:
        print(f"Disimpan di folder proyek: {project_path}")

if __name__ == "__main__":
    create_document()
